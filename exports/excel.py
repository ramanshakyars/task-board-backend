from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def export_tasks_to_excel(queryset) -> BytesIO:
 
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tasks"

    # Headr row stying
    headr_fill = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
    headr_font = Font(bold=True, color="FFFFFF")

    headers = [
        "ID", "Code", "Name", "Priority", "Status",
        "Description", "Due Date", "Created By", "Created At",
    ]

    for col_num, header in enumerate(headers, start=1):
        cel = ws.cell(row=1, column=col_num, value=header)
        cel.font = headr_font
        cel.fill = headr_fill
        cel.alignment = Alignment(horizontal="center")

    # Data ros
    for task in queryset:
        ws.append([
            task.id,
            task.code,
            task.name,
            task.get_priority_display(),
            task.get_status_display(),
            task.description,
            str(task.due_date) if task.due_date else "",
            task.created_by.username if task.created_by else "",
            task.created_at.strftime("%Y-%m-%d %H:%M") if task.created_at else "",
        ])

    # Auto-siz columns
    for col in ws.columns:
        max_lenght = max((len(str(cel.value or "")) for cel in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max(max_lenght + 4, 12)

    ouput = BytesIO()
    wb.save(ouput)
    ouput.seek(0)
    return ouput
